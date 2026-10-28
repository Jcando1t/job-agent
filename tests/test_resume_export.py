from io import BytesIO
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
from fastapi import HTTPException
from PIL import Image

import main
from app.services.resume_export import render_resume_docx
from app.services.resume_export import render_resume_pdf


class ResumeExportServiceTests(unittest.TestCase):
    def test_docx_export_preserves_headings_lists_and_bold_text(self) -> None:
        output = render_resume_docx(
            "AI 应用工程师简历",
            "# 张三\n## 项目经历\n- **负责内容：** 使用 Python\n- 结果：完成原型",
        )
        document = Document(BytesIO(output))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]

        self.assertIn("张三", paragraphs)
        self.assertIn("项目经历", paragraphs)
        self.assertIn("负责内容： 使用 Python", paragraphs)
        self.assertEqual(document.core_properties.title, "AI 应用工程师简历")

    def test_docx_export_sets_east_asian_font_for_chinese_glyphs(self) -> None:
        document = Document(BytesIO(render_resume_docx("测试简历", "# 张三\n中文内容")))

        for style_name in ("Normal", "Heading 1", "List Bullet"):
            style = document.styles[style_name]
            self.assertEqual(
                style._element.rPr.rFonts.get(qn("w:eastAsia")),
                "Microsoft YaHei",
            )

    def test_docx_export_uses_a4_page_size(self) -> None:
        document = Document(BytesIO(render_resume_docx("测试简历", "# 张三")))
        section = document.sections[0]

        self.assertAlmostEqual(section.page_width, Mm(210), delta=1000)
        self.assertAlmostEqual(section.page_height, Mm(297), delta=1000)

    def test_docx_export_distinguishes_name_and_section_headings(self) -> None:
        document = Document(BytesIO(render_resume_docx("测试简历", "# 张三\n## 项目经历")))

        self.assertEqual(document.styles["Title"].font.size.pt, 18)
        self.assertEqual(str(document.styles["Heading 2"].font.color.rgb), "087E69")

    def test_personal_and_contact_details_move_into_resume_header(self) -> None:
        content = "# 个人简历\n**基本信息**\n姜林峰 | 2004.07 | 应届生\n电话：13800000000 | 邮箱：a@example.com\n**求职意向**\nAI 应用开发工程师\n**教育背景**\n某大学"
        document = Document(BytesIO(render_resume_docx("个人简历", content)))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]

        self.assertIn("姜林峰", paragraphs)
        self.assertTrue(any("电话：13800000000" in paragraph for paragraph in paragraphs))
        self.assertTrue(any("求职意向：AI 应用开发工程师" in paragraph for paragraph in paragraphs))
        self.assertNotIn("基本信息", paragraphs)
        self.assertNotIn("求职意向", paragraphs)
        self.assertIn("教育背景", paragraphs)

    def test_pdf_export_is_one_page_and_keeps_chinese_text(self) -> None:
        from pypdf import PdfReader

        output = render_resume_pdf("测试简历", "# 张三\n## 项目经历\n- 使用 Python 完成项目")
        reader = PdfReader(BytesIO(output))
        self.assertEqual(len(reader.pages), 1)
        self.assertIn("张三", reader.pages[0].extract_text())

    def test_pdf_export_promotes_contact_block_to_header(self) -> None:
        from pypdf import PdfReader

        content = "# 个人简历\n**基本信息**\n姜林峰 | 2004.07 | 应届生\n电话：13800000000 | 邮箱：a@example.com\n**求职意向**\nAI 应用开发工程师\n**教育背景**\n某大学"
        text = PdfReader(BytesIO(render_resume_pdf("个人简历", content))).pages[0].extract_text()

        self.assertIn("姜林峰", text)
        self.assertIn("求职意向：AI 应用开发工程师", text)
        self.assertNotIn("基本信息", text)

    def test_photo_is_embedded_in_pdf_and_docx(self) -> None:
        photo = BytesIO()
        Image.new("RGB", (200, 280), color="navy").save(photo, format="PNG")
        photo_bytes = photo.getvalue()

        document = Document(BytesIO(render_resume_docx("测试简历", "# 张三\n## 项目经历", photo_bytes)))
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(0, 0).text, "张三")

        from pypdf import PdfReader

        pdf = PdfReader(BytesIO(render_resume_pdf("测试简历", "# 张三\n## 项目经历", photo_bytes)))
        self.assertEqual(len(pdf.pages), 1)
        self.assertTrue(pdf.pages[0].images)


class ResumeExportApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.resume = {
            "id": 5,
            "title": "AI 应用工程师简历",
            "content": "# 张三\n## 技能\n- Python",
        }

    def test_downloads_docx_with_safe_attachment_headers(self) -> None:
        with patch("main.get_resume", return_value=self.resume):
            response = main.download_resume(5)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.media_type,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("filename=\"resume-5.docx\"", response.headers["content-disposition"])
        self.assertIn("filename*=UTF-8''AI%20%E5%BA%94", response.headers["content-disposition"])
        Document(BytesIO(response.body))

    def test_downloads_markdown_without_changing_saved_content(self) -> None:
        with patch("main.get_resume", return_value=self.resume):
            response = main.download_resume(5, "md")

        self.assertTrue(response.media_type.startswith("text/markdown"))
        self.assertEqual(response.body.decode("utf-8"), self.resume["content"])

    def test_download_pdf_uses_the_selected_photo(self) -> None:
        resume = {**self.resume, "photo_image_id": None}
        photo = {"id": 9, "content": b"photo-bytes", "content_type": "image/png"}
        with patch("main.get_resume", return_value=resume), patch(
            "main.get_resume_image_source", return_value=photo
        ), patch("main.render_resume_pdf", return_value=b"pdf-bytes") as render:
            response = main.download_resume(5, "pdf", 9)

        self.assertEqual(response.body, b"pdf-bytes")
        render.assert_called_once_with("AI 应用工程师简历", self.resume["content"], b"photo-bytes")

    def test_rejects_missing_versions_and_unsupported_formats(self) -> None:
        with patch("main.get_resume", return_value=None):
            with self.assertRaises(HTTPException) as missing:
                main.download_resume(999)
        self.assertEqual(missing.exception.status_code, 404)

        with patch("main.get_resume", return_value=self.resume):
            with self.assertRaises(HTTPException) as unsupported:
                main.download_resume(5, "txt")
        self.assertEqual(unsupported.exception.status_code, 422)


if __name__ == "__main__":
    unittest.main()
