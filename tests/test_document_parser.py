import base64
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

import main
from app.services.document_parser import DocumentParseError, extract_resume_text
from app.schemas import ResumeFileInput


class DocumentParserTests(unittest.TestCase):
    def test_plain_text_is_decoded(self) -> None:
        self.assertEqual(extract_resume_text("resume.txt", "姓名：张三".encode("utf-8")), "姓名：张三")

    def test_legacy_chinese_text_encoding_is_decoded(self) -> None:
        self.assertEqual(extract_resume_text("resume.txt", "姓名：张三".encode("gb18030")), "姓名：张三")

    def test_unsupported_extension_is_rejected(self) -> None:
        with self.assertRaises(DocumentParseError):
            extract_resume_text("resume.png", b"image")

    def test_docx_text_is_extracted(self) -> None:
        stream = BytesIO()
        document = Document()
        document.add_paragraph("项目：求职工作流")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "技能"
        table.cell(0, 1).text = "Python"
        document.save(stream)
        extracted = extract_resume_text("resume.docx", stream.getvalue())
        self.assertIn("项目：求职工作流", extracted)
        self.assertIn("技能 | Python", extracted)

    def test_pdf_without_text_explains_scanned_document_limit(self) -> None:
        stream = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.write(stream)
        with self.assertRaisesRegex(DocumentParseError, "扫描图片"):
            extract_resume_text("resume.pdf", stream.getvalue())

    def test_upload_endpoint_extracts_text_and_stores_original_file(self) -> None:
        payload = ResumeFileInput(
            filename="C:\\fakepath\\resume.txt",
            content_base64=base64.b64encode("姓名：张三".encode("utf-8")).decode("ascii"),
        )
        with patch("main.create_resume_source", return_value={"id": 3, "content": "姓名：张三"}) as create:
            result = main.upload_resume_source(payload)
        self.assertEqual(result["id"], 3)
        self.assertEqual(create.call_args.kwargs["filename"], "resume.txt")
        self.assertEqual(create.call_args.kwargs["content"], "姓名：张三")
        Path("data/files").joinpath(create.call_args.kwargs["storage_path"]).unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
