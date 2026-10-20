"""读取用户主动上传的文本、PDF 和 Word 简历，输出可检索的纯文本。"""

from io import BytesIO
from pathlib import Path


class DocumentParseError(ValueError):
    """文件格式支持但内容无法读取。"""


def extract_resume_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".txt", ".md", ".text"}:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                # Windows 上创建的简历 TXT 常见 GBK/GB18030 编码。
                text = content.decode("gb18030")
        elif suffix == ".pdf":
            from pypdf import PdfReader

            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
        elif suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        else:
            raise DocumentParseError("暂不支持此文件格式，请上传 TXT、MD、PDF 或 DOCX。")
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError("文件无法读取，请确认文件没有损坏并重试。") from error

    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentParseError("文件中没有读取到文字。扫描图片请使用图片 OCR 后校对。")
    return text[:30000]
