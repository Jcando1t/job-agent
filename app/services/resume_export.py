"""将已保存的 Markdown 简历版本导出为常见文件格式。"""

from io import BytesIO
import html
import re
from pathlib import Path

from PIL import Image as PILImage, ImageOps
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.platypus.doctemplate import LayoutError


class ResumePdfExportError(ValueError):
    """简历无法安全地导出为一页 PDF。"""


_PDF_FONT = "STSong-Light"
_PDF_TTF_FONT = "ResumeYaHei"
_ACCENT = HexColor("#087E69")
_RULE = HexColor("#DCE4DF")


def render_resume_pdf(title: str, markdown: str, photo_content: bytes | None = None) -> bytes:
    """将审核后的 Markdown 简历导出为仅一页的 PDF。"""
    pdf_font = _register_pdf_font()

    styles = {
        "title": ParagraphStyle(
            "ResumeTitle", fontName=pdf_font, fontSize=20, leading=24,
            alignment=TA_LEFT, textColor=HexColor("#202C35"), spaceAfter=5,
        ),
        "section": ParagraphStyle(
            "ResumeSection", fontName=pdf_font, fontSize=11, leading=14,
            textColor=_ACCENT, spaceBefore=5, spaceAfter=1, keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "ResumeBody", fontName=pdf_font, fontSize=9.5, leading=12.5,
            textColor=HexColor("#34423C"), spaceAfter=2,
        ),
        "bullet": ParagraphStyle(
            "ResumeBullet", fontName=pdf_font, fontSize=9.5, leading=12.5,
            textColor=HexColor("#34423C"), leftIndent=10, firstLineIndent=-8,
            spaceAfter=2.5,
        ),
    }
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=11 * mm,
        bottomMargin=11 * mm,
        title=title,
    )
    story = []
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    header = _extract_resume_header(lines, title)
    photo = Image(BytesIO(_photo_thumbnail(photo_content)), width=25 * mm, height=31 * mm) if photo_content else None
    if photo:
        story.append(_pdf_title_row(_pdf_inline(header["name"]), photo, document.width, header["details"]))
    else:
        story.append(Paragraph(_pdf_inline(header["name"]), styles["title"]))
        story.extend(Paragraph(_pdf_inline(detail), styles["body"]) for detail in header["details"])
    story.append(HRFlowable(width="100%", thickness=0.8, color=_ACCENT, spaceBefore=0, spaceAfter=3))

    first_heading_seen = False
    for index, raw_line in enumerate(lines):
        if index in header["skip"]:
            continue
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            is_main_title = len(heading.group(1)) == 1 and not first_heading_seen
            story.append(Paragraph(_pdf_inline(heading.group(2)), styles["section"]))
            if is_main_title:
                story.append(HRFlowable(width="100%", thickness=0.8, color=_ACCENT, spaceBefore=0, spaceAfter=3))
            elif len(heading.group(1)) > 1:
                story.append(HRFlowable(width="100%", thickness=0.45, color=_RULE, spaceBefore=0, spaceAfter=3))
            first_heading_seen = True
        elif _standalone_bold_heading(line):
            story.append(Paragraph(_pdf_inline(_standalone_bold_heading(line)), styles["section"]))
            story.append(HRFlowable(width="100%", thickness=0.45, color=_RULE, spaceBefore=0, spaceAfter=3))
        elif line.startswith(("- ", "* ", "• ")):
            story.append(Paragraph(f'<font name="Helvetica" color="#087E69">-</font>&nbsp;&nbsp;{_pdf_inline(line[2:].strip())}', styles["bullet"]))
        elif re.match(r"^\d+[.、)]\s+", line):
            text = re.sub(r"^\d+[.、)]\s+", "", line)
            story.append(Paragraph(f'<font name="Helvetica" color="#087E69">-</font>&nbsp;&nbsp;{_pdf_inline(text)}', styles["bullet"]))
        else:
            story.append(Paragraph(_pdf_inline(line), styles["body"]))
    if not story:
        story.append(Paragraph(_pdf_inline(title), styles["title"]))

    try:
        document.build(story)
    except LayoutError as error:
        raise ResumePdfExportError("简历内容无法排入一页，请精简内容后再导出。") from error

    pdf = output.getvalue()
    if len(PdfReader(BytesIO(pdf)).pages) != 1:
        raise ResumePdfExportError("简历超过一页，请精简内容后再导出。")
    return pdf


def _pdf_inline(text: str) -> str:
    """Escape plain text and preserve basic Markdown emphasis."""
    parts = re.split(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`)", text)
    rendered = []
    for part in parts:
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            rendered.append(f"<b>{html.escape(part[2:-2])}</b>")
        elif part.startswith("`") and part.endswith("`"):
            rendered.append(html.escape(part[1:-1]))
        else:
            rendered.append(html.escape(part))
    return "".join(rendered)


def _photo_thumbnail(photo_content: bytes) -> bytes:
    """将用户照片裁切为统一的竖版头像尺寸，避免拉伸变形。"""
    image = ImageOps.exif_transpose(PILImage.open(BytesIO(photo_content))).convert("RGB")
    image = ImageOps.fit(image, (390, 480), method=PILImage.Resampling.LANCZOS)
    output = BytesIO()
    image.save(output, format="JPEG", quality=88)
    return output.getvalue()


def _pdf_title_row(text: str, photo: Image, width: float, details: list[str] | None = None) -> Table:
    title = Paragraph(text, ParagraphStyle("ResumePhotoTitle", fontName=_register_pdf_font(), fontSize=20, leading=24, textColor=HexColor("#202C35"), spaceAfter=3))
    left = [title, *(Paragraph(_pdf_inline(detail), ParagraphStyle("ResumeHeaderDetail", fontName=_register_pdf_font(), fontSize=8.5, leading=10, textColor=HexColor("#52635A"))) for detail in (details or []))]
    table = Table([[left, photo]], colWidths=[width - 27 * mm, 27 * mm])
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (1, 0), (1, 0), "RIGHT"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    return table


def _register_pdf_font() -> str:
    """Use a real local TrueType font so Latin and Chinese glyphs keep normal spacing."""
    if _PDF_TTF_FONT in pdfmetrics.getRegisteredFontNames():
        return _PDF_TTF_FONT
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    if font_path.is_file():
        try:
            pdfmetrics.registerFont(TTFont(_PDF_TTF_FONT, str(font_path), subfontIndex=0))
            return _PDF_TTF_FONT
        except (TypeError, ValueError):
            pass
    if _PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(_PDF_FONT))
    return _PDF_FONT


def _standalone_bold_heading(line: str) -> str | None:
    match = re.fullmatch(r"\*\*(.+?)\*\*", line.strip())
    return match.group(1).strip() if match else None


def _extract_resume_header(lines: list[str], fallback_title: str) -> dict:
    """Move name, contact details, and target role beside the optional photo."""
    skip: set[int] = set()
    main_heading = next(
        ((index, match.group(1).strip()) for index, line in enumerate(lines)
         if (match := re.match(r"^\s*#\s+(.+)$", line))),
        (None, fallback_title),
    )
    if main_heading[0] is not None:
        skip.add(main_heading[0])

    sections: dict[str, tuple[int, list[tuple[int, str]]]] = {}
    current_section: str | None = None
    for index, line in enumerate(lines):
        bold_heading = _standalone_bold_heading(line)
        markdown_heading = re.match(r"^\s*#{2,6}\s+(.+)$", line)
        label = bold_heading or (markdown_heading.group(1).strip() if markdown_heading else None)
        if label:
            normalized = label.rstrip("：:").strip()
            sections[normalized] = (index, [])
            current_section = normalized
            continue
        if current_section and line.strip():
            sections[current_section][1].append((index, line.strip()))

    info = next((value for key, value in sections.items() if key in {"基本信息", "个人信息"}), None)
    details: list[str] = []
    name = main_heading[1]
    if info:
        info_heading, info_lines = info
        skip.add(info_heading)
        skip.update(index for index, _ in info_lines)
        if info_lines:
            fields = [part.strip() for part in info_lines[0][1].split("|") if part.strip()]
            if fields:
                name = re.sub(r"^(姓名|名字)\s*[：:]\s*", "", fields[0])
                details.extend(fields[1:])
        if len(info_lines) > 1:
            details.append(info_lines[1][1])

    intent = next((value for key, value in sections.items() if key in {"求职意向", "目标岗位", "求职目标"}), None)
    if intent:
        intent_heading, intent_lines = intent
        skip.add(intent_heading)
        if intent_lines:
            skip.add(intent_lines[0][0])
            details.append(f"求职意向：{intent_lines[0][1]}")

    return {"name": name, "details": details, "skip": skip}


def render_resume_docx(title: str, markdown: str, photo_content: bytes | None = None) -> bytes:
    """将简历 Markdown 的标题和列表转换为可编辑的 Word 文档。"""
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    # Word uses separate font settings for Latin and East Asian characters.
    # Setting only font.name can leave Chinese glyphs to an unsupported theme font.
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )
        style.paragraph_format.line_spacing = 1.0

    normal = document.styles["Normal"]
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(52, 66, 60)
    normal.paragraph_format.space_after = Pt(0.5)
    normal.paragraph_format.line_spacing = 0.98
    for style_name in ("List Bullet", "List Number"):
        document.styles[style_name].font.size = Pt(9)
        document.styles[style_name].paragraph_format.left_indent = Inches(0.18)
        document.styles[style_name].paragraph_format.first_line_indent = Inches(-0.12)
    for style_name, size, before, after in (
        ("Title", 18, 0, 2),
        ("Heading 1", 11, 3, 1),
        ("Heading 2", 10.5, 2, 0.5),
        ("Heading 3", 10, 2, 0.5),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(8, 126, 105) if style_name.startswith("Heading") else RGBColor(32, 44, 53)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    document.core_properties.title = title

    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    first_heading_seen = False
    photo_added = False
    header = _extract_resume_header(lines, title)
    if photo_content:
        _add_docx_title_row(document, header["name"], photo_content, header["details"])
        photo_added = True
    else:
        paragraph = document.add_paragraph(style="Title")
        _add_inline_text(paragraph, header["name"])
        for detail in header["details"]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline_text(paragraph, detail)

    for index, raw_line in enumerate(lines):
        if index in header["skip"]:
            continue
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            is_main_title = len(heading.group(1)) == 1 and not first_heading_seen
            if photo_content and not photo_added and is_main_title:
                _add_docx_title_row(document, header["name"], photo_content, header["details"])
                photo_added = True
            elif is_main_title:
                paragraph = document.add_paragraph(style="Title")
                _add_inline_text(paragraph, header["name"])
                for detail in header["details"]:
                    paragraph = document.add_paragraph()
                    _add_inline_text(paragraph, detail)
            else:
                document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 3))
            first_heading_seen = True
        elif _standalone_bold_heading(line):
            document.add_heading(_standalone_bold_heading(line), level=2)
        elif line.startswith(("- ", "* ", "• ")):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_text(paragraph, line[2:].strip())
        elif re.match(r"^\d+[.、)]\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_text(paragraph, re.sub(r"^\d+[.、)]\s+", "", line))
        else:
            paragraph = document.add_paragraph()
            _add_inline_text(paragraph, line)

    if not first_heading_seen and not markdown.strip():
        document.add_heading(title, level=0)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _add_docx_title_row(document: Document, title: str, photo_content: bytes, details: list[str] | None = None) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.7)
    table.columns[1].width = Inches(0.9)
    left, right = table.rows[0].cells
    left.width = Inches(5.7)
    right.width = Inches(0.9)
    left.paragraphs[0].style = "Title"
    _add_inline_text(left.paragraphs[0], title)
    right.paragraphs[0].alignment = 2
    for detail in details or []:
        paragraph = left.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        _add_inline_text(paragraph, detail)
    right.paragraphs[0].add_run().add_picture(BytesIO(_photo_thumbnail(photo_content)), width=Mm(21), height=Mm(26))


def _add_inline_text(paragraph, text: str) -> None:
    """保留 Markdown 粗体和行内代码的可读文字及基础样式。"""
    token_pattern = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`)")
    for part in token_pattern.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(part)
